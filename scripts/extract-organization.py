"""Read the supplied organization workbook into review candidates; never invent HR identifiers/dates."""
import argparse, json, re
from pathlib import Path
import openpyxl

parser=argparse.ArgumentParser()
parser.add_argument('workbook',type=Path)
parser.add_argument('--out',type=Path,default=Path('data/imports/besakih-draft.json'))
args=parser.parse_args()
book=openpyxl.load_workbook(args.workbook,data_only=True)
main=book['Bagan Lengkap']
branches={'Bagan Sekretariat':'Sekretariat','Bagan SPI':'Satuan Pengawas Internal & Manajemen Risiko','Bagan Operasional':'Bidang Operasional & Pelayanan','Bagan Asset':'Bidang Pengelolaan Aset & Pengembangan Usaha','Bagan Lingkungan':'Bidang Pengelolaan Lingkungan & Keamanan'}
title_branches={}
for sheet_name,branch in branches.items():
    sheet=book[sheet_name]
    # Copying only labels under this branch avoids interpreting drawing connector lines as authority.
    for row in sheet:
        for cell in row:
            if cell.value is not None and cell.row>=23:
                value=str(cell.value).strip()
                if re.match(r'^(KEPALA URUSAN|KOORD\.|STAFF? |STAF |THL |WAKIL )',value,re.I): title_branches[value]=branch

def clean_name(value):
    value=re.sub(r'^\s*\d+\s*[.)]\s*','',value).strip()
    value=re.sub(r'\s*\(Pjs\.?\)\s*','',value,flags=re.I).strip()
    return re.sub(r'\s+',' ',value)

by_name={};positions={};skipped=[]
header_rows=[5,14,23,32,41,50]
for header_row in header_rows:
    for row in main.iter_rows(min_row=header_row,max_row=header_row):
        for cell in row:
            if cell.value is None: continue
            title=str(cell.value).strip()
            if title.startswith('BAGAN'): continue
            branch=title_branches.get(title,'Badan Pengelola Besakih')
            # Top-level functions are explicitly titled in the source.
            if header_row==14:
                if 'SEKRETARIS' in title: branch='Sekretariat'
                elif 'PENGAWAS' in title: branch=branches['Bagan SPI']
                elif 'OPERASIONAL' in title: branch=branches['Bagan Operasional']
                elif 'ASSET' in title or 'ASET' in title: branch=branches['Bagan Asset']
                elif 'LINGKUNGAN' in title: branch=branches['Bagan Lingkungan']
            positions[(branch,title)]={'name':title,'org_path':branch}
            name_rows=range(header_row+3,main.max_row+1) if header_row==50 else [header_row+3]
            for name_row in name_rows:
                source=main.cell(name_row,cell.column)
                if source.value is None: continue
                raw=str(source.value).strip()
                for part in raw.split(';'):
                    name=clean_name(part)
                    if not name or '?' in name or re.search(r'\bTBA\b',name,re.I):
                        skipped.append({'cell':source.coordinate,'value':part,'reason':'Nama belum ditetapkan'})
                        continue
                    key=name.casefold()
                    candidate=by_name.setdefault(key,{'full_name':name,'positions':[],'source_cells':[],'employee_no':None,'joined_on':None,'employment_status':None,'review_required':True})
                    if {'name':title,'org_path':branch} not in candidate['positions']:candidate['positions'].append({'name':title,'org_path':branch})
                    candidate['source_cells'].append(f'Bagan Lengkap!{source.coordinate}')

payload={'schema_version':1,'source_name':args.workbook.name,'source_sheet':'Bagan Lengkap','notes':'Hanya nama dan label fungsi dari bagan. Penempatan utama, status resmi, nomor pegawai dan tanggal mulai memerlukan tinjauan pemilik data. Tidak menetapkan kewenangan akses atau kelayakan pos keamanan.','org_units':['Badan Pengelola Besakih',*branches.values()],'positions':list(positions.values()),'candidates':list(by_name.values()),'unresolved':skipped}
args.out.parent.mkdir(parents=True,exist_ok=True)
args.out.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'candidates':len(by_name),'positions':len(positions),'unresolved':len(skipped),'output':str(args.out)},ensure_ascii=False))
